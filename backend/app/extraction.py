"""Text-layer extraction for accepted documents. Extraction reads only text
already present in a document's text layer -- no OCR, no LLM involvement.
The result is a document's full plain text, not field-level key/value
records; that remains a separate, future capability."""

import io

from docx import Document as DocxDocument
from pptx import Presentation
from pypdf import PdfReader
from sqlalchemy.orm import Session

from app.models import Document, DocumentExtraction
from app.storage import storage_client


class NoExtractableTextError(Exception):
    """Raised when extraction yields no non-whitespace text."""


def _extract_pdf_text(content: bytes) -> str:
    reader = PdfReader(io.BytesIO(content))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _extract_docx_text(content: bytes) -> str:
    document = DocxDocument(io.BytesIO(content))
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def _extract_pptx_text(content: bytes) -> str:
    presentation = Presentation(io.BytesIO(content))
    lines = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                lines.append(shape.text_frame.text)
    return "\n".join(lines)


def _extract_txt_text(content: bytes) -> str:
    try:
        return content.decode("utf-8")
    except UnicodeDecodeError:
        return content.decode("latin-1")


_EXTRACTORS = {
    "pdf": _extract_pdf_text,
    "docx": _extract_docx_text,
    "pptx": _extract_pptx_text,
    "txt": _extract_txt_text,
}


def extract_text(format: str, content: bytes) -> str:
    """Extract text-layer content for a supported format. Raises
    NoExtractableTextError if no non-whitespace text is found -- shared
    across formats so no extractor has to detect "no text" on its own."""
    extractor = _EXTRACTORS[format]
    text = extractor(content)
    if not text or not text.strip():
        raise NoExtractableTextError("No text layer was found in the document.")
    return text


def run_extraction(document: Document, db: Session) -> DocumentExtraction:
    """Extract text for `document` and persist the outcome as a
    DocumentExtraction row. Any failure -- no text layer, a corrupt file, or
    any other extractor exception -- is captured as a failed outcome with a
    reason rather than propagated, so it cannot affect the already-created
    Document record or stored file."""
    try:
        content = storage_client.get(document.storage_key)
        text = extract_text(document.format, content)
    except NoExtractableTextError as exc:
        extraction = DocumentExtraction(
            document_id=document.id,
            extracted_text=None,
            status="failed",
            failure_reason=str(exc),
        )
    except Exception as exc:  # noqa: BLE001 - any extractor/storage failure must be captured, not propagated
        extraction = DocumentExtraction(
            document_id=document.id,
            extracted_text=None,
            status="failed",
            failure_reason=f"Extraction failed: {exc}",
        )
    else:
        extraction = DocumentExtraction(
            document_id=document.id,
            extracted_text=text,
            status="succeeded",
            failure_reason=None,
        )

    db.add(extraction)
    db.commit()
    db.refresh(extraction)
    return extraction
