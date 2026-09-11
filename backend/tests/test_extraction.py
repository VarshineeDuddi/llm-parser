import io
import uuid

import pytest
from docx import Document as DocxDocument
from pptx import Presentation
from pypdf import PdfWriter
from sqlalchemy import select

from app import extraction
from app.extraction import NoExtractableTextError, extract_text, run_extraction
from app.models import Document, DocumentExtraction
from app.storage import storage_client


def _build_pdf_with_text(text: str) -> bytes:
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 200 200] /Contents 4 0 R "
        "/Resources << /Font << /F1 5 0 R >> >> >>",
        None,
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    stream = f"BT /F1 24 Tf 10 100 Td ({text}) Tj ET"
    objects[3] = f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream"

    body = "%PDF-1.4\n"
    offsets = []
    for i, obj in enumerate(objects, start=1):
        offsets.append(len(body.encode("latin-1")))
        body += f"{i} 0 obj\n{obj}\nendobj\n"

    xref_offset = len(body.encode("latin-1"))
    xref = f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n"
    for offset in offsets:
        xref += f"{offset:010d} 00000 n \n"
    trailer = (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_offset}\n%%EOF"
    )

    return (body + xref + trailer).encode("latin-1")


def _build_blank_pdf() -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def _build_docx_with_text(text: str) -> bytes:
    doc = DocxDocument()
    doc.add_paragraph(text)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _build_pptx_with_text(text: str) -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = text
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


PDF_WITH_TEXT = _build_pdf_with_text("Hello PDF")
BLANK_PDF = _build_blank_pdf()
DOCX_WITH_TEXT = _build_docx_with_text("Hello DOCX")
PPTX_WITH_TEXT = _build_pptx_with_text("Hello PPTX")


def _make_document(format: str, content: bytes) -> Document:
    document_id = uuid.uuid4()
    storage_key = f"documents/{document_id}"
    storage_client.put(storage_key, content)
    return Document(
        id=document_id,
        original_filename=f"sample.{format}",
        format=format,
        size_bytes=len(content),
        storage_key=storage_key,
        status="received",
    )


# --- Format-specific extractors (spec: Text-Layer Extraction for Supported Formats) ---


def test_extract_pdf_text_with_text_layer():
    text = extract_text("pdf", PDF_WITH_TEXT)
    assert "Hello PDF" in text


def test_extract_pdf_no_text_layer_raises():
    with pytest.raises(NoExtractableTextError):
        extract_text("pdf", BLANK_PDF)


def test_extract_docx_text():
    text = extract_text("docx", DOCX_WITH_TEXT)
    assert "Hello DOCX" in text


def test_extract_pptx_text():
    text = extract_text("pptx", PPTX_WITH_TEXT)
    assert "Hello PPTX" in text


def test_extract_txt_utf8():
    text = extract_text("txt", "héllo wörld".encode("utf-8"))
    assert text == "héllo wörld"


def test_extract_txt_latin1_fallback():
    content = "café".encode("latin-1")
    text = extract_text("txt", content)
    assert text == "café"


def test_extract_text_no_non_whitespace_raises():
    with pytest.raises(NoExtractableTextError):
        extract_text("txt", b"   \n\t  ")


# --- Extraction service (spec: Extraction Outcome Tracking, Extraction Failure Isolation) ---


@pytest.mark.parametrize(
    ("format", "content", "expected_text"),
    [
        ("pdf", PDF_WITH_TEXT, "Hello PDF"),
        ("docx", DOCX_WITH_TEXT, "Hello DOCX"),
        ("pptx", PPTX_WITH_TEXT, "Hello PPTX"),
        ("txt", b"hello txt", "hello txt"),
    ],
)
def test_run_extraction_creates_row_for_each_format(db_session, s3, format, content, expected_text):
    document = _make_document(format, content)
    db_session.add(document)
    db_session.commit()
    db_session.refresh(document)

    result = run_extraction(document, db_session)

    assert result.status == "succeeded"
    assert expected_text in result.extracted_text
    assert result.document_id == document.id

    stored = db_session.execute(
        select(DocumentExtraction).where(DocumentExtraction.document_id == document.id)
    ).scalar_one()
    assert stored.id == result.id


def test_run_extraction_captures_extractor_exception(db_session, s3, monkeypatch):
    document = _make_document("txt", b"hello")
    db_session.add(document)
    db_session.commit()
    db_session.refresh(document)

    def _boom(content: bytes) -> str:
        raise RuntimeError("boom")

    monkeypatch.setitem(extraction._EXTRACTORS, "txt", _boom)

    result = run_extraction(document, db_session)

    assert result.status == "failed"
    assert "boom" in result.failure_reason


# --- Upload integration (spec: Extraction Triggered After Upload, Extraction Failure Isolation) ---


def test_upload_triggers_extraction_synchronously(client, db_session):
    response = client.post(
        "/documents/upload",
        files={"file": ("sample.txt", b"hello world", "text/plain")},
    )
    assert response.status_code == 201
    body = response.json()
    assert body["extraction_status"] == "succeeded"
    assert body["extraction_failure_reason"] is None

    extraction_row = db_session.execute(
        select(DocumentExtraction).where(
            DocumentExtraction.document_id == uuid.UUID(body["id"])
        )
    ).scalar_one()
    assert extraction_row.status == "succeeded"


def test_upload_extraction_failure_does_not_affect_document(client, db_session):
    response = client.post(
        "/documents/upload",
        files={"file": ("sample.pdf", BLANK_PDF, "application/pdf")},
    )
    assert response.status_code == 201
    body = response.json()
    assert body["extraction_status"] == "failed"
    assert body["extraction_failure_reason"]

    document = db_session.get(Document, uuid.UUID(body["id"]))
    assert document is not None
    assert document.status == "received"
    assert storage_client.get(document.storage_key) == BLANK_PDF


# --- Extraction retrieval (spec: Extracted Text Persistence, Extraction Outcome Tracking) ---


def test_get_extraction_returns_text_for_succeeded(client):
    response = client.post(
        "/documents/upload",
        files={"file": ("sample.txt", b"retrievable text", "text/plain")},
    )
    document_id = response.json()["id"]

    get_response = client.get(f"/documents/{document_id}/extraction")

    assert get_response.status_code == 200
    body = get_response.json()
    assert body["status"] == "succeeded"
    assert body["extracted_text"] == "retrievable text"
    assert body["failure_reason"] is None


def test_get_extraction_returns_failure_reason_not_text(client):
    response = client.post(
        "/documents/upload",
        files={"file": ("sample.pdf", BLANK_PDF, "application/pdf")},
    )
    document_id = response.json()["id"]

    get_response = client.get(f"/documents/{document_id}/extraction")

    assert get_response.status_code == 200
    body = get_response.json()
    assert body["status"] == "failed"
    assert body["extracted_text"] is None
    assert body["failure_reason"]
